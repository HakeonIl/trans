# Gameloc — 게임 폴더를 넣으면 번역해 주는 도구
# Copyright (C) 2026  하진
#
# 이 프로그램은 자유 소프트웨어입니다. 자유 소프트웨어 재단이 펴낸 GNU 일반
# 공중 사용 허가서 3판 또는 그 이후 판의 조건에 따라 재배포하거나 고칠 수
# 있습니다.
#
# 이 프로그램은 쓸모가 있기를 바라며 배포하지만 어떠한 보증도 하지 않습니다.
# 자세한 내용은 GNU 일반 공중 사용 허가서를 보세요.
# 함께 받은 LICENSE 파일에 전문이 있습니다. <https://www.gnu.org/licenses/>

"""번역 워크북: 원문 한 줄 = 한 행, 언어 하나 = 한 열.

왕복에 중요한 규칙:

* A열의 ID 가 ``project.json`` 으로 돌아가는 유일한 열쇠입니다. 행을 정렬·필터·
  숨기는 건 자유지만 지우면 안 되고, ID 를 고쳐서도 안 됩니다.
* 언어 열은 **위치가 아니라 머리글 이름** 으로 찾습니다. 번역자가 아무 데나
  새 언어 열을 끼워 넣어도 찾아냅니다.
* 도구가 쓴 칸은 회색, 사람이 채울 칸은 노란색입니다.
"""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.protection import SheetProtection

from . import privacy as _privacy
from .model import Entry, Project

SHEET_MAIN = "번역"
SHEET_INFO = "정보"

FIXED_HEADERS = ["ID", "횟수", "위치", "원문"]
TAIL_HEADERS = ["검토", "메모"]

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="2F3B52")
LOCKED_FILL = PatternFill("solid", fgColor="F2F2F2")
INPUT_FILL = PatternFill("solid", fgColor="FFF9D6")
THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# 구글 시트의 번역 수식(``=GOOGLETRANSLATE``)을 미리 채울 수 있는 언어 열 이름.
_LANG_CODE = re.compile(r"^[a-z]{2,3}(-[A-Za-z]{2,4})?$")

# 구글 시트가 **계산을 끝내지 못한 칸**에 남기는 글자. 수식이 아직 계산 중일 때
# 파일로 내려받으면 번역 대신 이 글자가 그대로 저장됩니다. 이걸 번역으로 읽어서
# 게임에 넣으면 대사 자리에 ``Loading...`` 이 박힙니다.
_UNFINISHED = {"loading...", "loading…", "loading", "#error!", "#error", "#name?",
               "#n/a", "#value!", "#ref!", "#div/0!", "#num!", "#null!"}


def is_unfinished(value) -> bool:
    """번역이 끝나지 않은 칸인가 (``Loading...``·오류 표시·계산 안 된 번역 수식).

    번역문이 ``=`` 로 시작하는 일은 실제로 있어서(``==== 시작 ====``) ``=`` 로
    시작하면 다 막지 않고, **우리가 채운 번역 수식**만 미완성으로 봅니다.
    """
    if not isinstance(value, str):
        return False
    v = value.strip()
    return v.lower() in _UNFINISHED or v.upper().startswith("=GOOGLETRANSLATE(")


def unfinished_message(problems: list[str]) -> str:
    """끝나지 않은 칸이 있을 때 사용자에게 보여 줄 말. 없으면 빈 문자열."""
    if not problems:
        return ""
    return (f"번역이 끝나지 않은 칸 {len(problems):,}개는 넣지 않았습니다 "
            "(Loading... · 오류 표시 · 계산 안 된 수식). 구글 시트에서 번역 열을 "
            "복사해 [값만 붙여넣기] 한 뒤 다시 내려받으세요. "
            f"예: {', '.join(problems[:3])}")


def write_workbook(proj: Project, path: Path, *, max_rows: int | None = None,
                   google_formulas: bool = False, src_lang: str = "auto") -> Path:
    """``google_formulas`` 가 켜지면 **비어 있는** 번역 칸에 구글 시트 번역 수식을
    채웁니다(이미 번역이 있는 칸은 그대로). 시트에서 열면 바로 번역됩니다."""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_MAIN

    headers = FIXED_HEADERS + list(proj.languages) + TAIL_HEADERS
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(vertical="center", horizontal="center")
        cell.border = BORDER
    ws.row_dimensions[1].height = 22

    entries = proj.entries[:max_rows] if max_rows else proj.entries
    lang_start = len(FIXED_HEADERS) + 1

    for e in entries:
        loc = e.locations[0].label() if e.locations else ""
        if e.count > 1:
            loc += f" 외 {e.count - 1}곳"
        row = [e.id, e.count, loc, e.text]
        r = ws.max_row + 1                      # 이 행이 들어갈 행 번호
        for lg in proj.languages:
            have = e.translations.get(lg, "")
            if (google_formulas and not have and e.text.strip()
                    and _LANG_CODE.match(lg)):
                # 원문은 D열(4번째). 열 자리를 고정($D)해서 끌어 채워도 안 밀립니다.
                have = f'=GOOGLETRANSLATE($D{r},"{src_lang}","{lg}")'
            row.append(have)
        row += ["기계" if e.machine else "", e.note]
        ws.append(row)

    last = ws.max_row
    for r in range(2, last + 1):
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name=FONT, size=10)
            cell.border = BORDER
            is_input = (lang_start <= c < lang_start + len(proj.languages)
                        or c == len(headers))
            cell.fill = INPUT_FILL if is_input else LOCKED_FILL
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=c >= len(FIXED_HEADERS),
                horizontal="center" if c == 2 else "left",
            )
            cell.protection = Protection(locked=not is_input)

    widths = {1: 12, 2: 7, 3: 46, 4: 60}
    for i in range(len(proj.languages)):
        widths[lang_start + i] = 60
    widths[len(headers)] = 24
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w

    ws.freeze_panes = f"{get_column_letter(lang_start)}2"
    if last >= 2:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last}"
    # **시트를 잠그지 않습니다.**
    #
    # 예전에는 ID·원문 칸을 실수로 고치지 못하게 시트 보호를 걸었습니다.
    # 엑셀은 "이 칸만 잠금 해제" 를 이해하지만, **폴라리스 오피스는 보호된
    # 시트를 통째로 막습니다** — 번역을 붙여넣지도 못합니다.
    #
    #     보호된 시트는 편집할 수 없습니다
    #
    # 지키려던 것(ID 가 망가지는 것)보다 막아 버린 것(번역을 못 넣는 것)이
    # 훨씬 흔하고 나쁩니다. 잠금 대신 **회색 칠로 표시만** 하고, 어긋난 것은
    # 되읽을 때 걸러냅니다.
    ws.protection = SheetProtection(sheet=False)

    _write_info(wb, proj, shown=len(entries), google_formulas=google_formulas)
    wb.save(path)
    return path


def _write_info(wb: Workbook, proj: Project, shown: int,
                google_formulas: bool = False) -> None:
    ws = wb.create_sheet(SHEET_INFO)
    rows = []
    if google_formulas:
        rows += [
            ("구글 시트로 번역하기", ""),
            ("", "1. 이 파일을 구글 드라이브에 올리고 [구글 스프레드시트로 열기] 를 누르세요."),
            ("", "2. 번역 칸에 Loading... 이 하나도 안 보일 때까지 기다리세요. "
                 "줄이 많아 멈추면 나눠서 하세요."),
            ("", "3. 번역 열을 전체 선택 → 복사 → 같은 자리에 [값만 붙여넣기] (Ctrl+Shift+V). "
                 "이래야 수식이 사라지고 진짜 글자만 남습니다."),
            ("", "4. 파일 → 다운로드 → Microsoft Excel(.xlsx) 로 내려받아 작업 폴더의 "
                 "translation.xlsx 를 덮어쓰세요."),
            ("", "5. 번역기에서 [엑셀 ↑] 를 누르세요. 끝나지 않은 칸은 넣지 않고 알려 드립니다."),
            ("", ""),
        ]
    rows += [
        ("사용법", ""),
        ("", "노란색 칸에만 번역을 입력하세요. 회색 칸(ID·원문·위치)은 고치지 마세요."),
        ("", "ID 로 짝을 맞추므로 ID 가 바뀌면 그 줄의 번역이 안 들어갑니다."),
        ("", "행을 지우지 마세요. 번역하지 않은 행은 비워두면 원문이 그대로 유지됩니다."),
        ("", "언어를 추가하려면 헤더에 새 열을 만들고 언어 코드(예: ja, en)를 적으세요."),
        ("", "저장한 뒤 번역기에서 [엑셀 ↑] 를 누르면 읽어옵니다."),
        ("", ""),
        ("예시", ""),
        ("원문 예", "Press any key to continue"),
        ("번역 예", "아무 키나 누르세요"),
        ("", ""),
        ("게임 폴더", _privacy.scrub(proj.game_root)),
        ("엔진", f"{proj.engine} {proj.unity_version}".strip()),
        ("추출 시각", proj.generated_at),
        ("언어 열", ", ".join(proj.languages)),
        ("표시된 행", f"{shown} / {len(proj.entries)}"),
        ("", ""),
        ("--- 스캔 통계 ---", ""),
    ]
    rows += [(k, str(v)) for k, v in sorted(proj.stats.items())]
    for r in rows:
        ws.append(list(r))
    for row in ws.iter_rows():
        for cell in row:
            cell.font = Font(name=FONT, size=10, bold=cell.column == 1)
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column == 2)
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 90


def read_workbook(path: Path, problems: list[str] | None = None
                  ) -> tuple[dict[str, dict[str, str]], list[str]]:
    """``({entry_id: {lang: text}}, 찾은 언어들)``.

    ``Loading...``·오류 표시·계산 안 된 번역 수식이 든 칸은 **번역으로 읽지 않고
    건너뜁니다.** ``problems`` 리스트를 넘기면 건너뛴 칸을 ``"ID (언어)"`` 로 채워
    줍니다(사용자에게 알리는 용도).
    """
    wb = load_workbook(path, data_only=True, read_only=True)
    if SHEET_MAIN not in wb.sheetnames:
        raise ValueError(f"'{SHEET_MAIN}' 시트를 찾을 수 없습니다: {path}")
    ws = wb[SHEET_MAIN]
    # 같은 시트를 **수식 그대로**도 읽습니다. 값만 읽으면 "계산 결과가 없는 수식"과
    # "그냥 빈 칸"이 똑같이 None 으로 보여서, 계산이 안 끝난 칸을 놓칩니다.
    wb_f = load_workbook(path, data_only=False, read_only=True)
    ws_f = wb_f[SHEET_MAIN]

    rows = ws.iter_rows(values_only=True)
    rows_f = ws_f.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    next(rows_f, None)
    known = set(FIXED_HEADERS) | set(TAIL_HEADERS)
    lang_cols = {h: i for i, h in enumerate(header) if h and h not in known}
    try:
        id_col = header.index("ID")
    except ValueError as e:
        raise ValueError("헤더에 'ID' 열이 없습니다. 첫 행을 지우지 마세요.") from e

    out: dict[str, dict[str, str]] = {}
    bad: list[str] = []
    for row, frow in zip(rows, rows_f):
        if row is None or id_col >= len(row):
            continue
        eid = row[id_col]
        if not eid:
            continue
        vals = {}
        for lang, ci in lang_cols.items():
            v = row[ci] if ci < len(row) else None
            if v is None or str(v).strip() == "":
                fv = frow[ci] if frow is not None and ci < len(frow) else None
                if is_unfinished(fv):           # 수식만 있고 계산 결과가 없음
                    bad.append(f"{str(eid).strip()} ({lang})")
                continue
            if is_unfinished(v):                # Loading... · #ERROR! 등
                bad.append(f"{str(eid).strip()} ({lang})")
                continue
            vals[lang] = str(v)
        if vals:
            out[str(eid).strip()] = vals
    wb.close()
    wb_f.close()
    if problems is not None:
        problems.extend(bad)
    return out, list(lang_cols.keys())


def restore_newlines(source: str, translated: str) -> str:
    """구글 시트를 거치며 **두 배가 된 줄바꿈**을 되돌립니다.

    진짜 게임(Freedom)에서 원문의 ``\\r\\n`` 하나가 번역에서 ``\\n\\n`` 두 개가 되어
    돌아왔습니다(번역된 15,000여 개 중 5,806개). 대사창에 빈 줄이 하나씩 더 생깁니다.

    **딱 맞아떨어질 때만** 고칩니다. 원문이 ``\\r\\n`` 만 쓰고(맨 ``\\n`` 이나 ``\\r`` 이
    없고), 번역에는 ``\\r`` 이 없으며, 번역의 ``\\n`` 개수가 원문 ``\\r\\n`` 개수의 정확히
    두 배이고, 모든 ``\\n`` 이 짝을 이룰 때뿐입니다. 손으로 일부러 넣은 빈 줄이나
    번역가가 줄 수를 바꾼 경우는 아무것도 바꾸지 않습니다.
    """
    n = source.count("\r\n")
    if (n == 0 or source.count("\r") != n or source.count("\n") != n
            or "\r" in translated or translated.count("\n") != 2 * n):
        return translated
    fixed = translated.replace("\n\n", "\r\n")
    if fixed.count("\r\n") != n or fixed.count("\n") != n:
        return translated
    return fixed


def merge_into_project(proj: Project, edits: dict[str, dict[str, str]]) -> int:
    by_id = {e.id: e for e in proj.entries}
    changed = 0
    for eid, vals in edits.items():
        e: Entry | None = by_id.get(eid)
        if e is None:
            continue
        for lang, text in vals.items():
            text = restore_newlines(e.text, text)
            if e.translations.get(lang) != text:
                e.translations[lang] = text
                # 사람이 손으로 고친 것이므로 검토 대기 해제
                if lang in e.machine:
                    e.machine.remove(lang)
                changed += 1
            if lang not in proj.languages:
                proj.languages.append(lang)
    return changed
